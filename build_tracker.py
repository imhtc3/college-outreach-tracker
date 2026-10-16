"""Builds College_Outreach_Tracker.xlsx with sample (fictional) campaign data.

Sheets: Guide, Colleges, Outreach Log, Daily Metrics, Dashboard, Templates.
All totals and rates are Excel formulas, so the workbook updates as you log new entries.
"""
import random
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

random.seed(7)

FONT = "Arial"
NAVY = "1F3A5F"
LIGHT = "EAF0F7"
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
BAND_FILL = PatternFill("solid", fgColor=LIGHT)
thin = Side(style="thin", color="C9D3E0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

START = date(2026, 10, 1)
DAYS = 14
REPORT_DATE = START + timedelta(days=DAYS - 1)

wb = Workbook()


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 30


def style_body(ws, first, last, ncols):
    for r in range(first, last + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            if (r - first) % 2 == 1:
                cell.fill = BAND_FILL


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------------------------------------------------------- Colleges
cities = ["Delhi", "Noida", "Gurugram", "Ghaziabad", "Faridabad", "Greater Noida"]
types = ["Engineering", "Management", "Arts & Science", "Commerce"]
prefixes = ["Northgate", "Riverside", "Sunrise", "Lakeview", "Greenfield", "Crescent",
            "Silver Oak", "Maple", "Horizon", "Pinecrest", "Westbrook", "Eastwood",
            "Bluebell", "Summit", "Kingsway", "Oakridge", "Brightwater", "Highland",
            "Starlight", "Meadow", "Redwood", "Clearview", "Fairmont", "Goldcrest",
            "Hillside", "Ivory", "Juniper", "Lotus", "Marigold", "Neelkanth"]
suffix = {"Engineering": "Institute of Technology", "Management": "School of Management",
          "Arts & Science": "College of Arts & Science", "Commerce": "College of Commerce"}
first_names = ["Anil", "Sunita", "Rakesh", "Meena", "Vikas", "Pooja", "Sanjay", "Neha",
               "Rajiv", "Kavita", "Amit", "Ritu", "Deepak", "Shalini", "Manoj"]
last_names = ["Sharma", "Verma", "Gupta", "Singh", "Kumar", "Mehta", "Arora", "Jain",
              "Malhotra", "Bansal"]
channels = ["Call", "Email", "WhatsApp"]
statuses = ["Not contacted", "Contacted", "Interested", "Onboarded", "Not interested"]

colleges = []
for i, p in enumerate(prefixes, start=1):
    t = random.choice(types)
    colleges.append({
        "id": f"C{i:03d}",
        "name": f"{p} {suffix[t]}",
        "city": random.choice(cities),
        "type": t,
        "tpo": f"{random.choice(first_names)} {random.choice(last_names)} (sample)",
        "pref": random.choice(channels),
        "students": random.choice([300, 450, 600, 800, 1000, 1200, 1500]),
    })

# ---------------------------------------------------------------- Outreach log (synthetic)
outcomes_tpo = ["No answer", "Asked to email", "Interested", "Shared with students",
                "Not interested", "Follow-up requested"]
log = []
college_state = {c["id"]: "Not contacted" for c in colleges}
for d in range(DAYS):
    day = START + timedelta(days=d)
    if day.weekday() == 6:  # no outreach on Sundays
        continue
    for _ in range(random.randint(6, 10)):
        c = random.choice(colleges[:26])  # last 4 colleges stay "Not contacted"
        ch = random.choices(channels, weights=[4, 3, 3])[0]
        who = random.choices(["TPO", "Student group"], weights=[3, 1])[0]
        if who == "TPO":
            out = random.choices(outcomes_tpo, weights=[3, 2, 3, 2, 1, 2])[0]
        else:
            out = random.choice(["Message shared", "Queries resolved", "Registration help"])
        log.append([day, c["id"], ch, who, out])
        if out in ("Interested", "Shared with students"):
            college_state[c["id"]] = "Interested"
        elif out == "Not interested" and college_state[c["id"]] != "Interested":
            college_state[c["id"]] = "Not interested"
        elif college_state[c["id"]] == "Not contacted":
            college_state[c["id"]] = "Contacted"

# Promote some interested colleges to onboarded (sharing registration links)
interested = [k for k, v in college_state.items() if v == "Interested"]
for k in random.sample(interested, k=max(1, len(interested) // 2)):
    college_state[k] = "Onboarded"

ws = wb.active
ws.title = "Colleges"
heads = ["College ID", "College", "City", "Type", "TPO (contact)", "Preferred channel",
         "Status", "Students reached", "Registrations", "Applications",
         "Touches logged", "Last contact", "Next follow-up", "Follow-up due?"]
ws.append(heads)
style_header(ws, 1, len(heads))
for idx, c in enumerate(colleges, start=2):
    st = college_state[c["id"]]
    reached = regs = apps = 0
    if st == "Onboarded":
        reached = int(c["students"] * random.uniform(0.4, 0.8))
        regs = int(reached * random.uniform(0.18, 0.35))
        apps = int(regs * random.uniform(0.35, 0.6))
    elif st == "Interested":
        reached = int(c["students"] * random.uniform(0.1, 0.3))
        regs = int(reached * random.uniform(0.05, 0.15))
        apps = int(regs * random.uniform(0.2, 0.4))
    ws.append([c["id"], c["name"], c["city"], c["type"], c["tpo"], c["pref"], st,
               reached, regs, apps,
               f"=COUNTIF('Outreach Log'!$B:$B,A{idx})",
               f"=IF(K{idx}=0,\"\",MAXIFS_PLACEHOLDER)",
               None, None])
    # Last contact: latest log date for this college
    ws.cell(row=idx, column=12).value = (
        f"=IF(K{idx}=0,\"\",_xlfn.MAXIFS('Outreach Log'!$A:$A,'Outreach Log'!$B:$B,A{idx}))")
    # Next follow-up: 3 days after last contact for open statuses
    ws.cell(row=idx, column=13).value = (
        f"=IF(OR(G{idx}=\"Onboarded\",G{idx}=\"Not interested\",L{idx}=\"\"),\"\",L{idx}+Dashboard!$C$4)")
    ws.cell(row=idx, column=14).value = (
        f"=IF(M{idx}=\"\",\"\",IF(M{idx}<=Dashboard!$C$3,\"Yes\",\"No\"))")
last_col_row = len(colleges) + 1
style_body(ws, 2, last_col_row, len(heads))
for r in range(2, last_col_row + 1):
    for col in (12, 13):
        ws.cell(row=r, column=col).number_format = "dd-mmm-yy"
    for col in (7, 8, 9, 10):
        ws.cell(row=r, column=col).fill = INPUT_FILL
widths(ws, [10, 34, 13, 15, 24, 12, 14, 11, 12, 12, 10, 12, 13, 11])
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{last_col_row}"

dv_status = DataValidation(type="list", formula1='"' + ",".join(statuses) + '"', allow_blank=True)
ws.add_data_validation(dv_status)
dv_status.add(f"G2:G{last_col_row + 100}")
dv_ch = DataValidation(type="list", formula1='"Call,Email,WhatsApp"', allow_blank=True)
ws.add_data_validation(dv_ch)
dv_ch.add(f"F2:F{last_col_row + 100}")

red = PatternFill("solid", fgColor="F8D7DA")
ws.conditional_formatting.add(f"N2:N{last_col_row}",
                              FormulaRule(formula=['$N2="Yes"'], fill=red,
                                          font=Font(name=FONT, bold=True, color="9C0006")))
green = PatternFill("solid", fgColor="D4EDDA")
ws.conditional_formatting.add(f"G2:G{last_col_row}",
                              FormulaRule(formula=['$G2="Onboarded"'], fill=green))

# ---------------------------------------------------------------- Outreach Log sheet
wl = wb.create_sheet("Outreach Log")
lheads = ["Date", "College ID", "Channel", "Contacted", "Outcome", "College", "Positive?"]
wl.append(lheads)
style_header(wl, 1, len(lheads))
positive = ["Interested", "Shared with students", "Follow-up requested",
            "Queries resolved", "Registration help"]
pos_list = ",".join(f'"{p}"' for p in positive)
for i, row in enumerate(log, start=2):
    wl.append(row + [
        f"=IFERROR(INDEX(Colleges!$B:$B,MATCH(B{i},Colleges!$A:$A,0)),\"Unknown ID\")",
        f"=IF(OR(E{i}=\"Interested\",E{i}=\"Shared with students\",E{i}=\"Follow-up requested\","
        f"E{i}=\"Queries resolved\",E{i}=\"Registration help\"),1,0)",
    ])
last_log = len(log) + 1
style_body(wl, 2, last_log, len(lheads))
for r in range(2, last_log + 1):
    wl.cell(row=r, column=1).number_format = "dd-mmm-yy"
    for col in range(1, 6):
        wl.cell(row=r, column=col).fill = INPUT_FILL
widths(wl, [12, 11, 11, 14, 22, 36, 10])
wl.freeze_panes = "A2"
wl.auto_filter.ref = f"A1:G{last_log}"
dv2 = DataValidation(type="list", formula1='"Call,Email,WhatsApp"', allow_blank=True)
wl.add_data_validation(dv2)
dv2.add(f"C2:C{last_log + 500}")
dv3 = DataValidation(type="list", formula1='"TPO,Student group"', allow_blank=True)
wl.add_data_validation(dv3)
dv3.add(f"D2:D{last_log + 500}")

# ---------------------------------------------------------------- Daily Metrics
wm = wb.create_sheet("Daily Metrics")
mheads = ["Date", "Calls", "Emails", "WhatsApp", "Total touches", "Positive responses",
          "Response rate", "New registrations", "New applications"]
wm.append(mheads)
style_header(wm, 1, len(mheads))
for d in range(DAYS):
    r = d + 2
    day = START + timedelta(days=d)
    regs = 0 if day.weekday() == 6 else random.randint(15, 70) + d * 4
    apps = int(regs * random.uniform(0.3, 0.5))
    wm.append([day,
               f"=COUNTIFS('Outreach Log'!$A:$A,A{r},'Outreach Log'!$C:$C,\"Call\")",
               f"=COUNTIFS('Outreach Log'!$A:$A,A{r},'Outreach Log'!$C:$C,\"Email\")",
               f"=COUNTIFS('Outreach Log'!$A:$A,A{r},'Outreach Log'!$C:$C,\"WhatsApp\")",
               f"=SUM(B{r}:D{r})",
               f"=SUMIFS('Outreach Log'!$G:$G,'Outreach Log'!$A:$A,A{r})",
               f"=IF(E{r}=0,0,F{r}/E{r})",
               regs, apps])
last_m = DAYS + 1
tot = last_m + 1
wm.cell(row=tot, column=1, value="Total")
for col in range(2, 7):
    L = get_column_letter(col)
    wm.cell(row=tot, column=col, value=f"=SUM({L}2:{L}{last_m})")
wm.cell(row=tot, column=7, value=f"=IF(E{tot}=0,0,F{tot}/E{tot})")
for col in (8, 9):
    L = get_column_letter(col)
    wm.cell(row=tot, column=col, value=f"=SUM({L}2:{L}{last_m})")
style_body(wm, 2, tot, len(mheads))
for r in range(2, tot + 1):
    wm.cell(row=r, column=1).number_format = "dd-mmm-yy"
    wm.cell(row=r, column=7).number_format = "0.0%"
    if r <= last_m:
        for col in (8, 9):
            wm.cell(row=r, column=col).fill = INPUT_FILL
for col in range(1, 10):
    wm.cell(row=tot, column=col).font = Font(name=FONT, bold=True)
widths(wm, [12, 9, 9, 11, 12, 13, 12, 14, 14])
wm.freeze_panes = "B2"

lc = LineChart()
lc.title = "Daily registrations vs applications"
lc.y_axis.title = "Count"
lc.height, lc.width = 8, 18
lc.add_data(Reference(wm, min_col=8, min_row=1, max_col=9, max_row=last_m), titles_from_data=True)
lc.set_categories(Reference(wm, min_col=1, min_row=2, max_row=last_m))
lc.x_axis.number_format = "dd-mmm"
wm.add_chart(lc, "K2")

# ---------------------------------------------------------------- Dashboard
wd = wb.create_sheet("Dashboard", 0)
wd.sheet_view.showGridLines = False
wd["B1"] = "College Outreach Campaign: Dashboard"
wd["B1"].font = Font(name=FONT, size=16, bold=True, color=NAVY)
wd["B3"], wd["C3"] = "Report date", REPORT_DATE
wd["B4"], wd["C4"] = "Follow-up gap (days)", 3
wd["C3"].number_format = "dd-mmm-yy"
for c in ("C3", "C4"):
    wd[c].fill = INPUT_FILL
    wd[c].font = Font(name=FONT, color="0000FF", bold=True)
wd["D3"] = "Edit the yellow cells. Everything else updates on its own."
wd["D3"].font = Font(name=FONT, italic=True, color="666666", size=9)

C = f"Colleges!$G$2:$G${last_col_row}"
kpis = [
    ("Colleges in list", f"=COUNTA(Colleges!$A$2:$A${last_col_row})", "0"),
    ("Colleges contacted", f"=COUNTIF({C},\"<>Not contacted\")", "0"),
    ("Colleges onboarded", f"=COUNTIF({C},\"Onboarded\")", "0"),
    ("Total outreach touches", "='Daily Metrics'!E" + str(tot), "0"),
    ("Positive response rate", "='Daily Metrics'!G" + str(tot), "0.0%"),
    ("Students reached", f"=SUM(Colleges!$H$2:$H${last_col_row})", "#,##0"),
    ("Registrations", f"=SUM(Colleges!$I$2:$I${last_col_row})", "#,##0"),
    ("Applications", f"=SUM(Colleges!$J$2:$J${last_col_row})", "#,##0"),
    ("Registration rate (of reached)", "=IF(C11=0,0,C12/C11)", "0.0%"),
    ("Application rate (of registered)", "=IF(C12=0,0,C13/C12)", "0.0%"),
    ("Follow-ups due now", f"=COUNTIF(Colleges!$N$2:$N${last_col_row},\"Yes\")", "0"),
]
wd["B5"] = "Key numbers"
wd["B5"].font = Font(name=FONT, bold=True, color="FFFFFF")
wd["B5"].fill = HEAD_FILL
wd["C5"].fill = HEAD_FILL
for i, (label, f, fmt) in enumerate(kpis, start=6):
    wd.cell(row=i, column=2, value=label).font = Font(name=FONT)
    cell = wd.cell(row=i, column=3, value=f)
    cell.number_format = fmt
    cell.font = Font(name=FONT, bold=True)
    for col in (2, 3):
        wd.cell(row=i, column=col).border = BORDER

# Funnel
wd["E5"], wd["F5"] = "Campaign funnel", "Count"
wd["E6"], wd["F6"] = "Colleges contacted", "=C7"
wd["E7"], wd["F7"] = "Colleges interested or onboarded", f"=COUNTIF({C},\"Interested\")+C8"
wd["E8"], wd["F8"] = "Colleges onboarded", "=C8"
# Status breakdown
wd["E10"], wd["F10"] = "College status", "Colleges"
for i, s in enumerate(statuses, start=11):
    wd.cell(row=i, column=5, value=s)
    wd.cell(row=i, column=6, value=f"=COUNTIF({C},E{i})")
# Channel performance
wd["H5"], wd["I5"], wd["J5"], wd["K5"] = "Channel", "Touches", "Positive", "Response rate"
for i, ch in enumerate(channels, start=6):
    wd.cell(row=i, column=8, value=ch)
    wd.cell(row=i, column=9, value=f"=COUNTIF('Outreach Log'!$C:$C,H{i})")
    wd.cell(row=i, column=10, value=f"=SUMIFS('Outreach Log'!$G:$G,'Outreach Log'!$C:$C,H{i})")
    wd.cell(row=i, column=11, value=f"=IF(I{i}=0,0,J{i}/I{i})").number_format = "0.0%"
# City breakdown
wd["H10"], wd["I10"], wd["J10"] = "City", "Onboarded", "Registrations"
for i, city in enumerate(cities, start=11):
    wd.cell(row=i, column=8, value=city)
    wd.cell(row=i, column=9, value=f"=COUNTIFS(Colleges!$C:$C,H{i},Colleges!$G:$G,\"Onboarded\")")
    wd.cell(row=i, column=10, value=f"=SUMIFS(Colleges!$I:$I,Colleges!$C:$C,H{i})")

for (r1, c1, r2, c2) in [(5, 5, 8, 6), (10, 5, 15, 6), (5, 8, 8, 11), (10, 8, 16, 10)]:
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = wd.cell(row=r, column=c)
            cell.border = BORDER
            if r == r1:
                cell.font = Font(name=FONT, bold=True, color="FFFFFF")
                cell.fill = HEAD_FILL
            elif cell.font is None or not cell.font.bold:
                cell.font = Font(name=FONT)
widths(wd, [2, 30, 14, 2, 30, 10, 2, 14, 11, 13, 13])

bc = BarChart()
bc.type = "bar"
bc.title = "College funnel"
bc.legend = None
bc.height, bc.width = 7, 14
bc.add_data(Reference(wd, min_col=6, min_row=5, max_row=8), titles_from_data=True)
bc.set_categories(Reference(wd, min_col=5, min_row=6, max_row=8))
bc.y_axis.majorGridlines = None
wd.add_chart(bc, "B19")

cc = BarChart()
cc.title = "Response rate by channel"
cc.legend = None
cc.height, cc.width = 7, 12
cc.add_data(Reference(wd, min_col=11, min_row=5, max_row=8), titles_from_data=True)
cc.set_categories(Reference(wd, min_col=8, min_row=6, max_row=8))
cc.y_axis.number_format = "0%"
wd.add_chart(cc, "H19")

# ---------------------------------------------------------------- Templates
wt = wb.create_sheet("Templates")
templates = [
    ("TPO: first email",
     "Subject: Career opportunities for your students – quick registration\n\n"
     "Dear [TPO name],\n\nI hope you're doing well. I'm reaching out about [initiative name], "
     "a free career programme that connects students with internships and jobs from [type of companies].\n\n"
     "Students only need to register through this link: [link]. It takes about two minutes.\n\n"
     "Could you share it with your final- and pre-final-year students? I'm happy to send a short "
     "poster or join a quick call to answer questions.\n\nThank you,\n[Your name]\n[Phone]"),
    ("TPO: phone call script",
     "1. Introduce yourself and the initiative in one line.\n"
     "2. Ask: 'Is this a good time for two minutes?'\n"
     "3. Explain the benefit for students (free, quick sign-up, real openings).\n"
     "4. Ask how they prefer to share it: email, WhatsApp groups or a notice.\n"
     "5. Confirm next step and date; log the call outcome right after."),
    ("TPO: follow-up (after 3 days)",
     "Dear [TPO name], just following up on my note about [initiative name]. "
     "Registrations close on [date]. I've attached a ready-to-share poster for your students. "
     "Please let me know if anything is unclear."),
    ("Students: WhatsApp group message",
     "Hi everyone! [Initiative name] is open for registrations. Free to join, with internship "
     "and job openings from [companies]. Register in 2 minutes: [link]. Last date: [date]. "
     "Message me here if you get stuck."),
    ("Students: registration reminder",
     "Reminder: [initiative name] registrations close on [date]. If you've registered, "
     "don't forget to apply to openings that match you: [link]."),
]
wt.append(["Use", "Message"])
style_header(wt, 1, 2)
for i, (name, body) in enumerate(templates, start=2):
    wt.append([name, body])
    wt.cell(row=i, column=1).font = Font(name=FONT, bold=True)
    wt.cell(row=i, column=2).font = Font(name=FONT, size=10)
    wt.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    wt.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    wt.row_dimensions[i].height = 15 * (body.count("\n") + 3)
widths(wt, [26, 95])

# ---------------------------------------------------------------- Guide
wg = wb.create_sheet("Guide", 0)
wg.sheet_view.showGridLines = False
rows = [
    ("College Outreach Campaign Tracker", None),
    ("", None),
    ("What it does", "Tracks a college outreach campaign end to end: which colleges and TPOs "
     "were contacted, by which channel, what they said, when to follow up, and how many "
     "students registered and applied."),
    ("How to use", "1. Add colleges on the Colleges sheet.\n2. Log every call, email or WhatsApp "
     "message on Outreach Log (one row per touch).\n3. Update status, students reached, "
     "registrations and applications for each college.\n4. Enter daily registrations and "
     "applications on Daily Metrics.\n5. Check Dashboard for the funnel, channel response "
     "rates and follow-ups due."),
    ("Legend", "Yellow cells are inputs you type. White cells are formulas: don't overwrite them. "
     "Red 'Yes' on Colleges means a follow-up is due."),
    ("Follow-up rule", "Colleges that are Contacted or Interested get a follow-up date "
     "3 days after the last touch (change the gap on Dashboard)."),
    ("Data note", "All colleges, contacts and numbers here are fictional sample data for "
     "demonstration (1–14 Oct 2026). Replace them with your real campaign data."),
]
for i, (a, b) in enumerate(rows, start=1):
    wg.cell(row=i, column=2, value=a)
    if b:
        wg.cell(row=i, column=3, value=b)
        wg.cell(row=i, column=2).font = Font(name=FONT, bold=True, color=NAVY)
        wg.cell(row=i, column=3).font = Font(name=FONT)
        wg.cell(row=i, column=3).alignment = Alignment(wrap_text=True, vertical="top")
        wg.cell(row=i, column=2).alignment = Alignment(vertical="top")
        wg.row_dimensions[i].height = 15 * (b.count("\n") + 2) + (15 if len(b) > 90 else 0)
wg["B1"].font = Font(name=FONT, size=16, bold=True, color=NAVY)
widths(wg, [2, 18, 90])

wb.active = 1  # open on Dashboard
wb.save("College_Outreach_Tracker.xlsx")
print("rows logged:", len(log), "statuses:", {s: list(college_state.values()).count(s) for s in statuses})
